import docx
from docx import Document #导入库

path = r"C:\Users\lxr\Desktop\demo.docx"
document = Document(path)
tables = document.tables
table = tables[0]
bm = []


def main():
    for i in range(1, len(table.rows)):
        result = table.cell(i, 5).text
        bm.append(result)


if __name__ == '__main__':
    main()
